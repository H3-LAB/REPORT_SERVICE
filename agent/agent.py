"""Agente do REPORT SERVICE: corre na máquina que chega às bases de dados e gera os reports SQL e script.

Só faz pedidos HTTPS de saída para a app: pede trabalho, gera o ficheiro e devolve-o com o log. O envio aos
destinatários é feito pela app.

    REPORT_SERVICE_URL=https://report-service.exemplo AGENT_TOKEN=rta_... python agent.py
"""

import contextlib
import csv
import io
import json
import logging
import os
import re
import signal
import socket
import subprocess
import sys
import tempfile
import time
import zipfile
from datetime import date, datetime, time as dtime
from decimal import Decimal
from zoneinfo import ZoneInfo

import openpyxl
import psycopg
import requests
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

VERSION = os.getenv('AGENT_VERSION', 'dev')
REPORT_SERVICE_URL = os.getenv('REPORT_SERVICE_URL', '').rstrip('/')
AGENT_TOKEN = os.getenv('AGENT_TOKEN', '')
ALLOW_SCRIPTS = os.getenv('ALLOW_SCRIPTS', 'false').strip().lower() in ('1', 'true', 'yes', 'sim')
POLL_SECONDS = int(os.getenv('POLL_SECONDS', '15'))
QUERY_TIMEOUT = int(os.getenv('QUERY_TIMEOUT', '900'))     # segundos por query
SCRIPT_TIMEOUT = int(os.getenv('SCRIPT_TIMEOUT', '1800'))  # segundos por script
# CA_BUNDLE: caminho do certificado da CA interna; VERIFY_TLS=false só em testes
VERIFY = os.getenv('CA_BUNDLE') or os.getenv('VERIFY_TLS', 'true').strip().lower() not in ('0', 'false', 'no')
HEADER_FILL = PatternFill(start_color='B0E0E6', end_color='B0E0E6', fill_type='solid')
ALLOWED_EXTENSIONS = ('.xlsx', '.zip', '.csv', '.pdf', '.txt')

logger = logging.getLogger('reporting.agent')
stopping = False


class Log(list):
    def add(self, message, level='info'):
        self.append({'time': datetime.now().strftime('%H:%M:%S'), 'level': level, 'message': str(message)[:2000]})
        logger.log(logging.WARNING if level in ('warn', 'error') else logging.INFO, message)


# ---------- Túnel SSH ----------

SSH_TUNNEL_WAIT = 20  # segundos até a porta local do túnel abrir
SSH_KEY_TYPE = re.compile(r'(ssh|ecdsa|sk)-\S+')


def host_keys(text):
    """Linhas de known_hosts (com qualquer nome, até em hash) ou ``.pub`` do servidor → ``tunnel <tipo> <chave>``,
    para usar com ``HostKeyAlias=tunnel``: só a chave conta, não o nome com que foi copiada."""
    keys = []
    for line in text.splitlines():
        fields = line.split()
        found = next((i for i, field in enumerate(fields[:-1]) if SSH_KEY_TYPE.fullmatch(field)), None)
        if found is not None:
            keys.append(f'tunnel {fields[found]} {fields[found + 1]}')
    return '\n'.join(keys) + '\n'


@contextlib.contextmanager
def ssh_tunnel(info, log):
    """Abre ``ssh -L`` para a base da ligação só durante o job e devolve a ligação a apontar para o túnel.
    Sem ``ssh`` na ligação, devolve-a tal como está. A chave e o known_hosts ficam num diretório temporário."""
    ssh = info.get('ssh')
    if not ssh:
        yield info
        return
    with tempfile.TemporaryDirectory(prefix='tunnel-') as work:
        key, known_hosts = os.path.join(work, 'key'), os.path.join(work, 'known_hosts')
        for path, text in ((key, ssh['key']), (known_hosts, host_keys(ssh['known_hosts']))):
            with open(os.open(path, os.O_WRONLY | os.O_CREAT, 0o600), 'w') as f:
                f.write(text)
        with socket.socket() as probe:
            probe.bind(('127.0.0.1', 0))
            port = probe.getsockname()[1]
        command = ['ssh', '-N', '-F', '/dev/null', '-i', key, '-p', str(ssh['port']),
                   '-o', 'IdentitiesOnly=yes', '-o', 'BatchMode=yes', '-o', 'LogLevel=ERROR',
                   '-o', f'UserKnownHostsFile={known_hosts}', '-o', 'HostKeyAlias=tunnel',
                   '-o', 'StrictHostKeyChecking=yes',
                   '-o', 'ExitOnForwardFailure=yes', '-o', 'ConnectTimeout=15', '-o', 'ServerAliveInterval=30',
                   '-L', f"127.0.0.1:{port}:{info['host']}:{info['port']}", f"{ssh['user']}@{ssh['host']}"]
        log.add(f"Túnel SSH {ssh['user']}@{ssh['host']}:{ssh['port']} para {info['host']}:{info['port']}")
        process = subprocess.Popen(command, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                                   stderr=subprocess.PIPE, text=True)
        try:
            deadline = time.monotonic() + SSH_TUNNEL_WAIT
            while True:
                if process.poll() is not None:
                    error = process.stderr.read().strip()[-500:] or f'código {process.returncode}'
                    raise RuntimeError(f'O túnel SSH falhou: {error}')
                try:
                    socket.create_connection(('127.0.0.1', port), timeout=1).close()
                    break
                except OSError:
                    if time.monotonic() > deadline:
                        raise RuntimeError(f'O túnel SSH não abriu em {SSH_TUNNEL_WAIT} s.')
                    time.sleep(0.2)
            yield {**info, 'host': '127.0.0.1', 'port': port}
        finally:
            process.terminate()
            try:
                process.wait(5)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait()
            log.add('Túnel SSH fechado')


# ---------- SQL ----------

PARAM = re.compile(r'(?<![:\w]):(from|to)\b')


def bind(query):
    """``:from``/``:to`` → parâmetros do psycopg (os ``%`` da query passam a ``%%``)."""
    return PARAM.sub(lambda m: f'%({m[1]})s', query.replace('%', '%%'))


def connect(info):
    conn = psycopg.connect(host=info['host'], port=info['port'], dbname=info['database'], user=info['username'],
                           password=info['password'], connect_timeout=15, application_name='report-service-agent')
    conn.read_only = True  # BEGIN READ ONLY: funciona também atrás do pgbouncer
    return conn


def run_query(conn, query, period):
    if period is None and PARAM.search(query):
        raise RuntimeError('A query usa :from/:to, mas o report está "Sem período": escolha um período no report.')
    with conn.transaction(), conn.cursor() as cur:
        cur.execute(f'SET LOCAL statement_timeout = {QUERY_TIMEOUT * 1000}')
        if period is None:
            cur.execute(query)  # sem parâmetros: os % da query ficam como estão
        else:
            cur.execute(bind(query), period)
        if cur.description is None:
            return [], []
        return [column.name for column in cur.description], cur.fetchall()


def cell(value, tz):
    """Valor que o Excel/CSV aceita: datas sem fuso (no fuso do report), JSON e UUID como texto."""
    if isinstance(value, datetime):
        return value.astimezone(tz).replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, dtime):
        return value.replace(tzinfo=None)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, (memoryview, bytes)):
        return bytes(value).hex()
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub('', value)
    if value is None or isinstance(value, (bool, int, float, Decimal, date)):
        return value
    return str(value)


def build_xlsx(sheets, tz):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, columns, rows in sheets:
        ws = wb.create_sheet(title=name)
        ws.append(columns)
        for header in ws[1]:
            header.fill, header.font = HEADER_FILL, Font(bold=True)
        widths = [len(str(c)) for c in columns]
        for index, row in enumerate(rows):
            values = [cell(v, tz) for v in row]
            ws.append(values)
            if index < 200:
                widths = [max(w, len(str(v)) if v is not None else 0) for w, v in zip(widths, values)]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
        ws.freeze_panes = 'A2'
        if columns:
            ws.auto_filter.ref = ws.dimensions
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_csv(columns, rows, tz):
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=';')  # ; é o que o Excel em português abre sem importar
    writer.writerow(columns)
    for row in rows:
        writer.writerow(['' if v is None else v.isoformat(sep=' ') if isinstance(v, datetime) else v
                         for v in (cell(v, tz) for v in row)])
    return buffer.getvalue().encode('utf-8-sig')


def zip_files(files):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in files:
            archive.writestr(name, data)
    return buffer.getvalue()


def sql_job(job, log):
    tz = ZoneInfo(job['timezone'])
    period = job['period'] and {k: datetime.fromisoformat(v) for k, v in job['period'].items()}
    sheets, connections = [], {}
    # Fecha as ligações e depois os túneis, também em erro
    with contextlib.ExitStack() as stack:
        for item in job['queries']:
            info = item['connection']
            if info['name'] not in connections:
                log.add(f"Ligação {info['name']} ({info['username']}@{info['host']}:{info['port']}/{info['database']})")
                conn = connect(stack.enter_context(ssh_tunnel(info, log)))
                stack.callback(conn.close)
                connections[info['name']] = conn
            started = time.monotonic()
            columns, rows = run_query(connections[info['name']], item['query'], period)
            log.add(f"Folha \"{item['sheet']}\": {len(rows)} linhas em {time.monotonic() - started:.1f} s", 'ok')
            sheets.append((item['sheet'], columns, rows))
    stem = job['file_stem']
    if job['output'] == 'csv_zip':
        return stem + '.zip', zip_files([(f'{name}.csv', build_csv(c, r, tz)) for name, c, r in sheets])
    data = build_xlsx(sheets, tz)
    if job['output'] == 'xlsx_zip':
        return stem + '.zip', zip_files([(stem + '.xlsx', data)])
    return stem + '.xlsx', data


# ---------- Script ----------

def env_name(name):
    return re.sub(r'\W', '_', name).upper()


def script_env(job, out_dir, connections):
    # Ambiente limpo: o script não vê o AGENT_TOKEN nem as outras variáveis do agente
    env = {key: os.environ[key] for key in ('PATH', 'LANG', 'LC_ALL', 'TZ', 'HTTPS_PROXY', 'HTTP_PROXY', 'NO_PROXY',
                                            'https_proxy', 'http_proxy', 'no_proxy') if key in os.environ}
    env.update(HOME=out_dir, REPORT_OUTPUT_DIR=out_dir, REPORT_NAME=job['report'], REPORT_FILE_STEM=job['file_stem'],
               REPORT_FROM=(job['period'] or {}).get('from', ''), REPORT_TO=(job['period'] or {}).get('to', ''),
               REPORT_TIMEZONE=job['timezone'],
               PYTHONUNBUFFERED='1')
    for info in connections:
        prefix = f"DB_{env_name(info['name'])}_"
        env.update({prefix + 'HOST': info['host'], prefix + 'PORT': str(info['port']), prefix + 'NAME': info['database'],
                    prefix + 'USER': info['username'], prefix + 'PASSWORD': info['password'],
                    prefix + 'DSN': psycopg.conninfo.make_conninfo(
                        host=info['host'], port=info['port'], dbname=info['database'], user=info['username'],
                        password=info['password'])})
    return env


def script_job(job, log):
    if not ALLOW_SCRIPTS:
        raise RuntimeError('Este agente não corre scripts (ALLOW_SCRIPTS=false no compose do agente).')
    with tempfile.TemporaryDirectory(prefix='reporting-') as work:
        out_dir = os.path.join(work, 'output')
        os.mkdir(out_dir)
        lang = job['script']['lang']
        path = os.path.join(work, 'script.py' if lang == 'python' else 'script.sh')
        with open(path, 'w') as f:
            f.write(job['script']['code'])
        command = [sys.executable, path] if lang == 'python' else ['bash', path]
        log.add(f'Script {lang} (limite {SCRIPT_TIMEOUT} s)')
        # Os túneis SSH das ligações ficam abertos só enquanto o script corre
        with contextlib.ExitStack() as stack:
            connections = [stack.enter_context(ssh_tunnel(info, log)) for info in job.get('connections', [])]
            try:
                done = subprocess.run(command, cwd=work, env=script_env(job, out_dir, connections), capture_output=True,
                                      text=True, timeout=SCRIPT_TIMEOUT)
            except subprocess.TimeoutExpired:
                raise RuntimeError(f'O script passou do limite de {SCRIPT_TIMEOUT} s.')
        for line in done.stdout.splitlines()[-200:]:
            log.add(line)
        for line in done.stderr.splitlines()[-100:]:
            log.add(line, 'warn')
        if done.returncode != 0:
            raise RuntimeError(f'O script terminou com o código {done.returncode}.')
        files = sorted(name for name in os.listdir(out_dir) if os.path.isfile(os.path.join(out_dir, name)))
        if not files:
            raise RuntimeError('O script não escreveu nenhum ficheiro em $REPORT_OUTPUT_DIR.')
        if len(files) == 1 and files[0].lower().endswith(ALLOWED_EXTENSIONS):
            with open(os.path.join(out_dir, files[0]), 'rb') as f:
                return files[0], f.read()
        contents = []
        for name in files:
            with open(os.path.join(out_dir, name), 'rb') as f:
                contents.append((name, f.read()))
        return job['file_stem'] + '.zip', zip_files(contents)


def execute(job):
    """``(ok, nome, dados, log)`` de um job; os erros ficam no log."""
    log = Log()
    try:
        name, data = (sql_job if job['kind'] == 'sql' else script_job)(job, log)
        log.add(f'Gerado {name} ({max(1, len(data) // 1024)} KB)', 'ok')
        return True, name, data, log
    except Exception as exc:  # o erro vai para a execução na app
        logger.exception('Job %s falhou', job.get('run_id'))
        log.add(f'Erro: {exc}'[:1000], 'error')
        return False, None, None, log


# ---------- Ciclo ----------

def main():
    logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
    if not REPORT_SERVICE_URL or not AGENT_TOKEN:
        sys.exit('Defina REPORT_SERVICE_URL e AGENT_TOKEN.')

    def stop(*_):
        global stopping
        stopping = True
    signal.signal(signal.SIGTERM, stop)
    signal.signal(signal.SIGINT, stop)

    session = requests.Session()
    session.headers['Authorization'] = f'Bearer {AGENT_TOKEN}'
    session.verify = VERIFY
    # Num contentor o hostname é o ID: AGENT_HOSTNAME dá o nome da máquina
    info = {'hostname': os.getenv('AGENT_HOSTNAME') or socket.gethostname(), 'version': VERSION, 'allow_scripts': ALLOW_SCRIPTS}
    logger.info('Agente %s em %s (scripts %s), a perguntar a %s de %s em %s s', VERSION, info['hostname'],
                'permitidos' if ALLOW_SCRIPTS else 'desligados', REPORT_SERVICE_URL, POLL_SECONDS, POLL_SECONDS)
    while not stopping:
        try:
            # Sem seguir redirecionamentos: um 3xx é um REPORT_SERVICE_URL errado e perdia o pedido
            response = session.post(f'{REPORT_SERVICE_URL}/api/agente/trabalho', json=info, timeout=30,
                                    allow_redirects=False)
            if response.status_code == 204:
                time.sleep(POLL_SECONDS)
                continue
            if response.is_redirect:
                logger.error('O servidor redirecionou para %s: REPORT_SERVICE_URL não aponta para a API dos agentes '
                             '(porta do agent-api, sem prefixo).', response.headers.get('Location'))
                time.sleep(60)
                continue
            if response.status_code == 401:
                logger.error('Token recusado: gere um novo em Administração → Agentes.')
                time.sleep(60)
                continue
            response.raise_for_status()
            job = response.json()
        except (requests.RequestException, ValueError) as exc:
            logger.warning('Sem ligação à app: %s', exc)
            time.sleep(POLL_SECONDS)
            continue
        logger.info('Execução #%s: %s (%s)', job['run_id'], job['report'], job['kind'])
        ok, name, data, log = execute(job)
        for attempt in range(5):
            try:
                session.post(f"{REPORT_SERVICE_URL}/api/agente/execucoes/{job['run_id']}",
                             data={'ok': '1' if ok else '0', 'log': json.dumps(log)},
                             files={'file': (name, data)} if ok else None, timeout=300,
                             allow_redirects=False).raise_for_status()
                break
            except requests.RequestException as exc:
                logger.warning('Falhou a entrega do resultado (%s/5): %s', attempt + 1, exc)
                time.sleep(10 * (attempt + 1))
        # Logo a seguir pergunta de novo: pode haver mais trabalho na fila


if __name__ == '__main__':
    main()
